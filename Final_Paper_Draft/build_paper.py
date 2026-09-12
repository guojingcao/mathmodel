from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SRC = Path(__file__).resolve().with_name("B题论文草稿.md")
OUT = Path(__file__).resolve().with_name("B题论文草稿.docx")


def set_east_asia(run, font_name: str) -> None:
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), font_name)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])
    set_east_asia(run, "Times New Roman")
    run.font.size = Pt(9)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.first_line_indent = Pt(21)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = "黑体"
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    title.paragraph_format.space_after = Pt(16)
    title_ppr = title._element.get_or_add_pPr()
    border = title_ppr.find(qn("w:pBdr"))
    if border is not None:
        title_ppr.remove(border)

    heading_specs = {
        "Heading 1": ("黑体", 14, 12, 6),
        "Heading 2": ("黑体", 12, 9, 4),
        "Heading 3": ("楷体", 11, 6, 3),
    }
    for name, (font, size, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = font
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True


def add_inline_markdown(paragraph, text: str, *, font_size: float = 10.5) -> None:
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            set_east_asia(run, "宋体")
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_east_asia(run, "Consolas")
            run.font.size = Pt(max(8.5, font_size - 1))
        else:
            run = paragraph.add_run(part)
            set_east_asia(run, "宋体")
        run.font.size = Pt(font_size)


def add_equation(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text.strip())
    set_east_asia(run, "Cambria Math")
    run.font.size = Pt(10.5)
    run.italic = True


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_repeat_table_header(table.rows[0])
    for i, row_data in enumerate(rows):
        row = table.rows[i]
        for j in range(cols):
            cell = row.cells[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            text = row_data[j] if j < len(row_data) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (i == 0 or j == 0 or len(text) < 18) else WD_ALIGN_PARAGRAPH.LEFT
            add_inline_markdown(p, text, font_size=8.5)
            if i == 0:
                set_cell_shading(cell, "D9EAF7")
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0, 0, 0)
            elif i % 2 == 0:
                set_cell_shading(cell, "F7FAFC")
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.first_line_indent = Pt(0)


def add_figure(doc: Document, rel_path: str, alt: str) -> None:
    img_path = (SRC.parent / rel_path).resolve()
    if not img_path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"[{alt}：图片待插入]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(img_path), width=Inches(5.55))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Pt(0)
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(6)
    cap.paragraph_format.keep_with_next = False
    cr = cap.add_run(alt)
    set_east_asia(cr, "宋体")
    cr.font.size = Pt(9)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def build() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.3)
    sec.right_margin = Cm(2.3)
    sec.header_distance = Cm(1.0)
    sec.footer_distance = Cm(1.0)
    add_page_number(sec.footer.paragraphs[0])
    configure_styles(doc)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_lines: list[str] = []
    first_h1 = True
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.left_indent = Pt(12)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(6)
                for k, code_line in enumerate(code_lines):
                    if k:
                        p.add_run().add_break()
                    run = p.add_run(code_line)
                    set_east_asia(run, "Consolas")
                    run.font.size = Pt(8.5)
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(raw)
            i += 1
            continue
        if not line:
            i += 1
            continue
        if line.startswith(">"):
            i += 1
            continue
        if line.startswith("# "):
            text = line[2:].strip()
            if first_h1:
                p = doc.add_paragraph(style="Title")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Pt(0)
                run = p.add_run(text)
                set_east_asia(run, "黑体")
                run.font.size = Pt(18)
                run.bold = True
                first_h1 = False
            else:
                doc.add_heading(text, level=1)
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_heading(line[3:].strip(), level=1)
            p.paragraph_format.first_line_indent = Pt(0)
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_heading(line[4:].strip(), level=2)
            p.paragraph_format.first_line_indent = Pt(0)
            i += 1
            continue
        image_match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line)
        if image_match:
            add_figure(doc, image_match.group(2), image_match.group(1))
            i += 1
            continue
        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        if line.startswith("$$"):
            eq = line[2:]
            if line.endswith("$$") and len(line) > 4:
                eq = line[2:-2]
                i += 1
            else:
                i += 1
                eq_lines = [eq]
                while i < len(lines) and not lines[i].strip().endswith("$$"):
                    eq_lines.append(lines[i].strip())
                    i += 1
                if i < len(lines):
                    eq_lines.append(lines[i].strip()[:-2])
                    i += 1
                eq = " ".join(eq_lines)
            add_equation(doc, eq)
            continue
        ordered = re.match(r"^(\d+)\.\s+(.*)$", line)
        bullet = re.match(r"^[-*]\s+(.*)$", line)
        if ordered or bullet:
            text = ordered.group(2) if ordered else bullet.group(1)
            p = doc.add_paragraph(style="List Bullet" if bullet else None)
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.left_indent = Pt(21 if bullet else 0)
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(2)
            add_inline_markdown(p, f"{ordered.group(1)}. {text}" if ordered else text)
            i += 1
            continue
        p = doc.add_paragraph()
        if line.startswith("**关键词：**"):
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_after = Pt(8)
        add_inline_markdown(p, line)
        i += 1

    props = doc.core_properties
    props.title = "有界测向误差下干扰源的确定性区域定位与机器人覆盖清除"
    props.subject = "2026年全国大学生数学建模竞赛B题论文草稿"
    props.author = "参赛队 202604004013"
    props.keywords = "有界测向误差; 可行区域; 观测几何; 确定性覆盖; 路径调度"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
